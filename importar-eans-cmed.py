#!/usr/bin/env python3
import argparse
import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
BULAS = ROOT / "bulas-medicamentos"
LAB_TOKENS = {
    "ache", "biolab", "cimed", "donaduzzi", "ems", "eurofarma", "germed",
    "glenmark", "medley", "neoquimica", "novaquimica", "pharlab", "prati",
    "ranbaxy", "sandoz", "teuto", "uniao", "uniaoquimica",
}
FORM_CODES = (
    (re.compile(r"\b(COM|COMP)\b"), "comprimido"),
    (re.compile(r"\b(CAP|CAPS)\b"), "capsula"),
    (re.compile(r"\bXPE\b"), "xarope"),
    (re.compile(r"\b(GOT|GTS)\b"), "gotas"),
    (re.compile(r"\b(SUS|SUSP)\b.*\bOR\b"), "suspensao oral"),
    (re.compile(r"\bSOL\b.*\bOR\b"), "solucao oral"),
    (re.compile(r"\b(CREM|CREME)\b"), "creme"),
    (re.compile(r"\b(POM|POMADA)\b"), "pomada"),
    (re.compile(r"\b(INJ|AMP)\b"), "injetavel"),
    (re.compile(r"\bCOL\b"), "colirio"),
    (re.compile(r"\bSPR\b"), "spray"),
    (re.compile(r"\bSACHE\b"), "sache"),
)
SALT_PREFIXES = (
    "bromidrato de ", "bissulfato de ", "besilato de ", "betaciclodextrina de ",
    "citrato de ", "cloridrato de ", "dicloridrato de ", "dimesilato de ",
    "fumarato de ", "hemifumarato de ", "hemitartarato de ", "maleato de ",
    "mesilato de ", "nitrato de ", "oxalato de ", "succinato de ",
    "sulfato de ", "tartarato de ",
)
SALT_SUFFIXES = {
    "anidra", "anidro", "calcica", "calcico", "di hidratada", "di hidratado",
    "diidratada", "diidratado", "dissodica", "dissodico", "hemihidratada",
    "hemihidratado", "hidratada", "hidratado", "monoidratada", "monoidratado",
    "potassica", "potassico", "sodica", "sodico",
}
FORM_WORDS = {
    "ampola", "ampolas", "capsula", "capsulas", "colirio", "comprimido",
    "comprimidos", "creme", "gel", "gotas", "injetavel", "pomada", "sache",
    "saches", "solucao", "suspensao", "oral", "spray", "xarope",
}


def normalize(value):
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def canonical_name(value):
    tokens = normalize(value).split()
    while len(tokens) > 1 and (tokens[-1] in LAB_TOKENS or re.fullmatch(r"\d+(mg|mcg|g|ml)", tokens[-1])):
        tokens.pop()
    return " ".join(tokens)


def ingredient_core(value):
    """Return a conservative salt-free key for a single active ingredient."""
    raw = str(value or "")
    if ";" in raw or "+" in raw:
        return ""
    key = canonical_name(raw)
    changed = True
    while changed:
        changed = False
        for prefix in SALT_PREFIXES:
            if key.startswith(prefix) and len(key) > len(prefix) + 2:
                key = key[len(prefix):].strip()
                changed = True
                break
    tokens = key.split()
    while tokens and tokens[-1] in SALT_SUFFIXES:
        tokens.pop()
    return " ".join(tokens)


def medicine_name_variants(value):
    raw = str(value or "")
    variants = {canonical_name(raw)}
    without_strength = re.sub(
        r"\b\d+(?:[,.]\d+)?\s*(?:mg|mcg|g|ui)(?:\s*/\s*(?:ml|g))?\b",
        " ",
        raw,
        flags=re.I,
    )
    tokens = canonical_name(without_strength).split()
    while tokens and (tokens[-1] in FORM_WORDS or tokens[-1] in LAB_TOKENS):
        tokens.pop()
    if tokens:
        variants.add(" ".join(tokens))
    for variant in tuple(variants):
        core = ingredient_core(variant)
        if core:
            variants.add(core)
    return {variant for variant in variants if len(variant) >= 3}


def active_ingredient_variants(value):
    exact = canonical_name(value)
    variants = {exact} if exact else set()
    core = ingredient_core(value)
    if core:
        variants.add(core)
    return variants


def slug_for(medicine):
    return Path(str(medicine.get("leafletPath") or "")).stem.lower()


def load_javascript_json(path, prefix):
    source = path.read_text(encoding="utf-8")
    source = re.sub(rf"^{re.escape(prefix)}\s*=\s*", "", source).rstrip(";\n ")
    return json.loads(source)


def extract_ean(value):
    digits = re.sub(r"\D", "", str(value or "").split(".")[0])
    return digits if len(digits) == 13 else ""


def presentation_details(value):
    raw = str(value or "")
    normalized = normalize(raw)
    strengths = re.findall(r"\d+(?:[,.]\d+)?\s*(?:mg|mcg|g|ui)(?:\s*/\s*(?:ml|g))?", raw, re.I)
    strength = " + ".join(re.sub(r"\s+", " ", item).strip() for item in strengths)
    form = next((label for pattern, label in FORM_CODES if pattern.search(raw.upper())), "")
    return strength, form, normalized


def presentation_score(medicine, strength, form):
    presentations = medicine.get("presentations") or []
    if not presentations:
        return 1
    wanted_strength = normalize(strength).replace(" ", "")
    wanted_form = normalize(form)
    best = 0
    for presentation in presentations:
        text = normalize(" ".join(str(presentation.get(key) or "") for key in ("strength", "form", "label")))
        compact = text.replace(" ", "")
        strength_match = not wanted_strength or wanted_strength in compact or compact in wanted_strength
        form_match = not wanted_form or wanted_form in text
        if strength_match and form_match:
            best = max(best, 3)
        elif strength_match:
            best = max(best, 2)
    return best


def build_name_index(medicines):
    index = {}
    for medicine in medicines:
        names = set()
        names.update(medicine_name_variants(medicine.get("name")))
        names.update(medicine_name_variants(slug_for(medicine).replace("-", " ")))
        for name in names:
            index.setdefault(name, []).append(medicine)
    return index


def find_candidate(name_index, product, substance, strength, form):
    ranked = []
    # The active ingredient is the canonical destination; the commercial name
    # remains searchable metadata for barcode and text searches.
    sources = [
        (active_ingredient_variants(substance), 140),
        (medicine_name_variants(product), 110),
    ]
    for keys, name_score in sources:
        for key in keys:
            for medicine in name_index.get(key, []):
                exact_bonus = int(key == canonical_name(substance)) * 10
                score = name_score + exact_bonus + presentation_score(medicine, strength, form)
                lab_penalty = int(slug_for(medicine).split("-")[-1] in LAB_TOKENS)
                ranked.append((score, -lab_penalty, slug_for(medicine), medicine))
    if not ranked:
        return None
    ranked.sort(key=lambda item: (-item[0], -item[1], item[2]))
    return ranked[0][3]


def main():
    parser = argparse.ArgumentParser(description="Associa EANs da farmacia a bulas usando a planilha oficial CMED/Anvisa.")
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--all-market", action="store_true", help="Importa todos os EANs da CMED que tenham uma bula compatível.")
    args = parser.parse_args()

    inventory_text = (BULAS / "EAN BULAS.rtf").read_text(encoding="utf-8", errors="ignore")
    inventory_eans = set(re.findall(r"\b\d{13}\b", inventory_text))
    medicines = [
        *load_javascript_json(BULAS / "medicine-library-raia.generated.js", "window.RAIA_MEDICINE_LIBRARY"),
        *load_javascript_json(BULAS / "medicine-library-xlsx.generated.js", "window.XLSX_MEDICINE_LIBRARY"),
    ]
    existing = load_javascript_json(BULAS / "ean-medicine-aliases.generated.js", "window.EAN_MEDICINE_ALIASES")
    name_index = build_name_index(medicines)

    workbook = load_workbook(args.xlsx, read_only=True, data_only=True)
    sheet = workbook.active
    official_rows = {}
    for row in sheet.iter_rows(min_row=43, values_only=True):
        for position in (5, 6, 7):
            ean = extract_ean(row[position])
            if ean and (args.all_market or ean in inventory_eans):
                official_rows[ean] = {
                    "substance": str(row[0] or "").strip(),
                    "laboratory": str(row[2] or "").strip(),
                    "product": str(row[8] or "").strip(),
                    "presentation": str(row[9] or "").strip(),
                }

    official_aliases = {}
    coverage = {}
    for ean, item in official_rows.items():
        strength, form, _ = presentation_details(item["presentation"])
        candidate = find_candidate(name_index, item["product"], item["substance"], strength, form)
        if not candidate:
            continue
        official_aliases[ean] = {
            "ean": ean,
            "candidateSlug": slug_for(candidate),
            "strength": strength,
            "form": form,
            "sourceName": f'{item["product"]} - {item["presentation"]} - {item["laboratory"]}',
            "sourceUrl": "https://www.gov.br/anvisa/pt-br/assuntos/medicamentos/cmed/precos",
            "source": "CMED/Anvisa",
            "activeIngredient": item["substance"],
            "commercialName": item["product"],
            "laboratory": item["laboratory"],
        }
        principle = canonical_name(item["substance"])
        summary = coverage.setdefault(principle, {
            "activeIngredient": item["substance"],
            "candidateSlug": slug_for(candidate),
            "commercialNames": set(),
            "eans": [],
        })
        summary["commercialNames"].add(item["product"])
        summary["eans"].append({
            "ean": ean,
            "commercialName": item["product"],
            "laboratory": item["laboratory"],
            "strength": strength,
            "form": form,
        })

    aliases = {**existing, **official_aliases}
    aliases = dict(sorted(aliases.items()))
    unmatched = sorted(inventory_eans - set(aliases))
    (BULAS / "ean-medicine-aliases.generated.js").write_text(
        f"window.EAN_MEDICINE_ALIASES = {json.dumps(aliases, ensure_ascii=False, separators=(',', ':'))};\n",
        encoding="utf-8",
    )
    (BULAS / "EANs-sem-associacao-confiavel.txt").write_text(
        "\n".join(unmatched) + ("\n" if unmatched else ""),
        encoding="utf-8",
    )
    coverage_report = []
    for item in coverage.values():
        item["commercialNames"] = sorted(item["commercialNames"])
        item["eans"] = sorted(item["eans"], key=lambda entry: entry["ean"])
        coverage_report.append(item)
    coverage_report.sort(key=lambda item: normalize(item["activeIngredient"]))
    (BULAS / "associacoes-ean-por-principio-ativo.json").write_text(
        json.dumps(coverage_report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({
        "inventory": len(inventory_eans),
        "marketRows": len(official_rows),
        "foundInCmed": len(official_rows),
        "associatedFromCmed": len(official_aliases),
        "totalAssociated": len(aliases),
        "activeIngredients": len(coverage_report),
        "commercialNames": sum(len(item["commercialNames"]) for item in coverage_report),
        "unmatched": len(unmatched),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
