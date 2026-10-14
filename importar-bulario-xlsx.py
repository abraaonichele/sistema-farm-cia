#!/usr/bin/env python3
import argparse
import html
import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parent
BULAS = ROOT / "bulas-medicamentos"
PDF_DIR = BULAS / "bulario-xlsx-pdfs"
LIBRARY_PATH = BULAS / "medicine-library-xlsx.generated.js"
RAIA_LIBRARY_PATH = BULAS / "medicine-library-raia.generated.js"

SECTION_PATTERNS = {
    "Como funciona": re.compile(r"(?:mecanismo\s+de\s+a[cç][aã]o\s+)?como\s+funciona\s*\?", re.I),
    "Como usar": re.compile(r"(?:posologia\s+)?como\s+usar\s*\?", re.I),
    "Contraindicações": re.compile(r"contraindica[cç][oõ]es(?:\s+quando\s+n[aã]o\s+devo\s+usar\s*\?)?", re.I),
    "Interação medicamentosa": re.compile(r"intera[cç][oõ](?:es)?\s+medicamentosas?", re.I),
}

ALL_HEADINGS = re.compile(
    r"(?:indica[cç][aã]o\s+para\s+que\s+serve\s*\?"
    r"|(?:mecanismo\s+de\s+a[cç][aã]o\s+)?como\s+funciona\s*\?"
    r"|(?:posologia\s+)?como\s+usar\s*\?"
    r"|contraindica[cç][oõ]es(?:\s+quando\s+n[aã]o\s+devo\s+usar\s*\?)?"
    r"|intera[cç][oõ](?:es)?\s+medicamentosas?"
    r"|efeitos?\s+colaterais(?:\s+quais\s+os\s+males\s+que\s+pode\s+me\s+causar\s*\?)?"
    r"|advert[eê]ncias(?:\s+e\s+precau[cç][oõ]es)?"
    r"|superdosagem(?:\s+o\s+que\s+fazer\s+se\s+usar\s+uma\s+quantidade\s+maior\s+do\s+que\s+a\s+indicada\s*\?)?"
    r"|armazenamento(?:\s+onde\s+como\s+e\s+por\s+quanto\s+tempo\s+posso\s+guardar\s*\?)?"
    r"|composi[cç][aã]o)",
    re.I,
)


def normalize(value):
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def slugify(value):
    return normalize(value).replace(" ", "-") or "medicamento"


def clean_text(value):
    text = unicodedata.normalize("NFC", str(value or ""))
    text = text.replace("\u200b", " ").replace("\ufeff", " ")
    text = re.sub(r"\bpublicidade\b", " ", text, flags=re.I)
    text = re.sub(r"\s+", " ", text)
    return text.strip(" .\t\r\n")


def title_case_sentences(value):
    text = clean_text(value)
    if not text:
        return ""
    return text[0].upper() + text[1:]


def extract_section(text, target_pattern):
    target = target_pattern.search(text)
    if not target:
        return ""
    next_heading = next((match for match in ALL_HEADINGS.finditer(text, target.end()) if match.start() > target.start()), None)
    content = text[target.end():next_heading.start() if next_heading else len(text)]
    return title_case_sentences(content)


def load_javascript_json(path, prefix):
    source = path.read_text(encoding="utf-8")
    source = re.sub(rf"^{re.escape(prefix)}\s*=\s*", "", source).rstrip(";\n ")
    return json.loads(source)


def unique_rows(sheet):
    rows_by_name = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        name = clean_text(row[2])
        if not name:
            continue
        item = {
            "name": name,
            "composition": title_case_sentences(row[1]),
            "text": clean_text(row[4]),
        }
        key = normalize(name)
        if key not in rows_by_name or len(item["text"]) > len(rows_by_name[key]["text"]):
            rows_by_name[key] = item
    return rows_by_name


def split_paragraphs(value, sentence_limit=5):
    sentences = re.split(r"(?<=[.!?])\s+(?=[a-záàâãéêíóôõúç0-9])", clean_text(value), flags=re.I)
    return [" ".join(sentences[index:index + sentence_limit]) for index in range(0, len(sentences), sentence_limit)]


def draw_page_number(canvas, document):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColorRGB(0.35, 0.4, 0.5)
    canvas.drawRightString(A4[0] - 18 * mm, 12 * mm, f"Página {document.page}")
    canvas.restoreState()


def build_pdf(path, medicine, composition, full_text):
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "MedicineTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=20, leading=24, textColor="#173F78", alignment=TA_CENTER, spaceAfter=10,
    )
    heading = ParagraphStyle(
        "SectionHeading", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=12, leading=15, textColor="#173F78", spaceBefore=10, spaceAfter=5,
    )
    body = ParagraphStyle(
        "Body", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=9.5, leading=13.5, textColor="#26354A", spaceAfter=6,
    )
    note = ParagraphStyle(
        "Note", parent=body, fontSize=8.5, leading=12, textColor="#667085", alignment=TA_CENTER,
    )
    document = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm, title=medicine["name"], author="Farmácia Superpreço",
    )
    story = [
        Paragraph(html.escape(medicine["name"]), title),
        Paragraph("Conteúdo importado do arquivo documentos_bulario.xlsx", note),
        Spacer(1, 5 * mm),
    ]
    if composition:
        story.append(Paragraph("Composição", heading))
        for paragraph in split_paragraphs(composition):
            story.append(Paragraph(html.escape(paragraph), body))
    story.append(Paragraph("Texto da bula", heading))
    for paragraph in split_paragraphs(full_text):
        story.append(Paragraph(html.escape(paragraph), body))
    document.build(story, onFirstPage=draw_page_number, onLaterPages=draw_page_number)


def main():
    parser = argparse.ArgumentParser(description="Importa bulas de documentos_bulario.xlsx para o sistema.")
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--skip-pdfs", action="store_true")
    args = parser.parse_args()

    workbook = load_workbook(args.xlsx, read_only=True, data_only=True)
    rows = unique_rows(workbook.active)
    current_medicines = load_javascript_json(RAIA_LIBRARY_PATH, "window.RAIA_MEDICINE_LIBRARY")
    current_names = {normalize(medicine.get("name")) for medicine in current_medicines}
    new_rows = [item for key, item in sorted(rows.items()) if key not in current_names]

    PDF_DIR.mkdir(parents=True, exist_ok=True)
    used_slugs = set()
    medicines = []
    missing_topics = {title: 0 for title in SECTION_PATTERNS}
    for item in new_rows:
        base_slug = slugify(item["name"])
        slug = base_slug
        suffix = 2
        while slug in used_slugs:
            slug = f"{base_slug}-{suffix}"
            suffix += 1
        used_slugs.add(slug)

        sections = []
        for title, pattern in SECTION_PATTERNS.items():
            content = extract_section(item["text"], pattern)
            if not content:
                missing_topics[title] += 1
                content = "Este tópico não foi identificado no texto fornecido para esta bula."
            sections.append({"title": title, "content": content})

        medicine = {
            "name": item["name"],
            "ean": "",
            "aliases": normalize(item["name"]),
            "source": "documentos_bulario.xlsx",
            "sourceLeaflet": "",
            "leafletPath": f"./bulas-medicamentos/bulario-xlsx-pdfs/{slug}.pdf",
            "presentations": [],
            "sections": sections,
        }
        medicines.append(medicine)
        if not args.skip_pdfs:
            build_pdf(PDF_DIR / f"{slug}.pdf", medicine, item["composition"], item["text"])

    LIBRARY_PATH.write_text(
        f"window.XLSX_MEDICINE_LIBRARY = {json.dumps(medicines, ensure_ascii=False, separators=(',', ':'))};\n",
        encoding="utf-8",
    )
    print(json.dumps({
        "spreadsheetRows": workbook.active.max_row - 1,
        "uniqueMedicines": len(rows),
        "alreadyRegistered": len(rows) - len(new_rows),
        "newMedicines": len(medicines),
        "pdfsCreated": 0 if args.skip_pdfs else len(medicines),
        "missingTopics": missing_topics,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
